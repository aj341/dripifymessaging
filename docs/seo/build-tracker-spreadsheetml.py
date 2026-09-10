import openpyxl, re, datetime as dt, warnings, html
warnings.filterwarnings('ignore')
SRC='designbees-aeo-seo-tracker-2026-09-10.xlsx'
wb=openpyxl.load_workbook(SRC)

def col_to_num(s):
    n=0
    for ch in s: n=n*26+(ord(ch)-64)
    return n

ref_re=re.compile(r"(?<![A-Za-z0-9_])(\$?)([A-Z]{1,3})(\$?)(\d{1,7})(?![A-Za-z0-9_(])")
def a1_to_r1c1(f, row, col):
    # protect quoted strings
    parts=[]; out=[]; i=0
    tokens=re.split(r'("(?:[^"]|"")*")', f)
    for t in tokens:
        if t.startswith('"') and t.endswith('"') and len(t)>=2:
            out.append(t); continue
        def rep(m):
            ac,c,ar,r=m.group(1),m.group(2),m.group(3),int(m.group(4))
            cn=col_to_num(c)
            cs=('C%d'%cn) if ac else ('C' if cn==col else 'C[%+d]'%(cn-col))
            rs=('R%d'%r) if ar else ('R' if r==row else 'R[%+d]'%(r-row))
            return rs+cs
        out.append(ref_re.sub(rep,t))
    return ''.join(out)

STYLES = '''<Styles>
<Style ss:ID="Default" ss:Name="Normal"><Alignment ss:Vertical="Top" ss:WrapText="1"/><Font ss:FontName="Arial" ss:Size="10"/></Style>
<Style ss:ID="tt"><Font ss:FontName="Arial" ss:Size="14" ss:Bold="1"/></Style>
<Style ss:ID="sub"><Font ss:FontName="Arial" ss:Size="10" ss:Italic="1" ss:Color="#555555"/><Alignment ss:Vertical="Top" ss:WrapText="1"/></Style>
<Style ss:ID="hd"><Font ss:FontName="Arial" ss:Size="10" ss:Bold="1" ss:Color="#FFFFFF"/><Interior ss:Color="#1F2A37" ss:Pattern="Solid"/><Alignment ss:Vertical="Center" ss:WrapText="1"/></Style>
<Style ss:ID="b"><Font ss:FontName="Arial" ss:Size="10" ss:Bold="1"/></Style>
<Style ss:ID="dt"><Font ss:FontName="Arial" ss:Size="10"/><NumberFormat ss:Format="yyyy\\-mm\\-dd"/></Style>
<Style ss:ID="pc"><Font ss:FontName="Arial" ss:Size="10"/><NumberFormat ss:Format="0%"/></Style>
</Styles>'''

def esc(s):
    return str(s).replace('&','&amp;').replace('<','&lt;').replace('>','&gt;').replace('"','&quot;')

out=['<?xml version="1.0"?>','<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" xmlns:o="urn:schemas-microsoft-com:office:office" xmlns:x="urn:schemas-microsoft-com:office:excel" xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">',STYLES]

for name in wb.sheetnames:
    ws=wb[name]
    maxr=ws.max_row; maxc=ws.max_column
    merges={}
    for m in ws.merged_cells.ranges:
        merges[(m.min_row,m.min_col)]=(m.max_row-m.min_row, m.max_col-m.min_col)
    out.append('<Worksheet ss:Name="%s">'%esc(name))
    cols=[]
    for ci in range(1,maxc+1):
        letter=openpyxl.utils.get_column_letter(ci)
        d=ws.column_dimensions.get(letter)
        if d is not None and d.width:
            cols.append('<Column ss:Index="%d" ss:Width="%d"/>'%(ci,round(d.width*6)))
    out.append('<Table>'+''.join(cols))
    prev_row=0
    for r in range(1,maxr+1):
        cells=[]
        last=0
        for c in range(1,maxc+1):
            cell=ws.cell(row=r,column=c)
            v=cell.value
            if v is None and (r,c) not in merges: continue
            attrs=''
            if c!=last+1: attrs+=' ss:Index="%d"'%c
            st=None
            if r==1: st='tt'
            elif r==2: st='sub'
            elif isinstance(v,str) and v in ('Metric','Value','Priority','Tasks','Baseline, 10 Sep 2026 (actual)','Source','How to use') and name=='Summary': st='b'
            fill = cell.fill
            if fill is not None and fill.fgColor is not None and fill.fgColor.rgb in ('001F2A37','FF1F2A37'): st='hd'
            if st is None:
                if cell.number_format and 'yy' in cell.number_format: st='dt'
                elif cell.number_format=='0%': st='pc'
            if st: attrs+=' ss:StyleID="%s"'%st
            if (r,c) in merges:
                dr,dc=merges[(r,c)]
                if dr: attrs+=' ss:MergeDown="%d"'%dr
                if dc: attrs+=' ss:MergeAcross="%d"'%dc
            if isinstance(v,str) and v.startswith('='):
                attrs+=' ss:Formula="%s"'%esc(a1_to_r1c1(v,r,c))
                cells.append('<Cell%s/>'%attrs)
            elif v is None:
                cells.append('<Cell%s/>'%attrs)
            elif isinstance(v,(dt.datetime,dt.date)):
                d2=v if isinstance(v,dt.datetime) else dt.datetime(v.year,v.month,v.day)
                cells.append('<Cell%s><Data ss:Type="DateTime">%s</Data></Cell>'%(attrs,d2.strftime('%Y-%m-%dT00:00:00.000')))
            elif isinstance(v,bool):
                cells.append('<Cell%s><Data ss:Type="Boolean">%d</Data></Cell>'%(attrs,int(v)))
            elif isinstance(v,(int,float)):
                cells.append('<Cell%s><Data ss:Type="Number">%s</Data></Cell>'%(attrs,('%g'%v) if isinstance(v,float) else str(v)))
            else:
                cells.append('<Cell%s><Data>%s</Data></Cell>'%(attrs,esc(v)))
            last=c
        if cells:
            h=ws.row_dimensions.get(r)
            ra=''
            while prev_row+1<r:
                out.append('<Row/>'); prev_row+=1
            if h is not None and h.height: ra+=' ss:AutoFitHeight="0" ss:Height="%g"'%h.height
            out.append('<Row%s>%s</Row>'%(ra,''.join(cells))); prev_row=r
    out.append('</Table>')
    fr = 4 if name in ('Tasks','Sources','Evergreen Pages','Tracking') else 0
    if fr:
        out.append('<WorksheetOptions xmlns="urn:schemas-microsoft-com:office:excel"><FreezePanes/><FrozenNoSplit/><SplitHorizontal>%d</SplitHorizontal><TopRowBottomPane>%d</TopRowBottomPane><ActivePane>2</ActivePane></WorksheetOptions>'%(fr,fr))
    out.append('</Worksheet>')
out.append('</Workbook>')
xml='\n'.join(out)
open('tracker.xml','w',encoding='utf-8').write(xml)
print('bytes',len(xml.encode()))
